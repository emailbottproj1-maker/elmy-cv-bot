"""
Demo: build a sample company database + Excel export to show the client.

Creates:
  - data/companies_db.sqlite   (queryable DB: name, email, language, sector, city, source)
  - data/companies_sample_list.xlsx  (ready to load into the bot)

This is a DEMO/starter dataset (public-style sample rows) to illustrate the
deliverable. The real campaign list is built/expanded as agreed with the client.
"""
import os
import sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DB = os.path.join(ROOT, "data", "companies_db.sqlite")
XLSX = os.path.join(ROOT, "data", "companies_sample_list.xlsx")

# Sample starter rows (illustrative). Sectors aligned to the 11 canonical
# codes recognised by the web UI: Healthcare, Marketing, Technology, Finance,
# Engineering, Legal, Retail, Education, HR, Hospitality, Logistics.
COMPANIES = [
    # --- Healthcare ---
    ("Dr. Sulaiman Al Habib", "hr@drhabib-example.com", "ar", "Healthcare", "Riyadh"),
    ("Mouwasat", "careers@mouwasat-example.com", "ar", "Healthcare", "Dammam"),
    ("Bupa Arabia", "hr@bupa-example.com", "en", "Healthcare", "Jeddah"),
    # --- Marketing ---
    ("J. Walter Thompson KSA", "jobs@jwt-example.com", "en", "Marketing", "Riyadh"),
    ("Lucidya", "careers@lucidya-example.com", "en", "Marketing", "Riyadh"),
    ("FP7McCann", "hr@fp7-example.com", "en", "Marketing", "Jeddah"),
    # --- Technology ---
    ("Elm", "jobs@elm-example.com", "ar", "Technology", "Riyadh"),
    ("Foodics", "hr@foodics-example.com", "en", "Technology", "Riyadh"),
    ("Unifonic", "careers@unifonic-example.com", "en", "Technology", "Riyadh"),
    # --- Finance ---
    ("Alinma Bank", "careers@alinma-example.com", "ar", "Finance", "Riyadh"),
    ("Riyad Bank", "recruitment@riyadbank-example.com", "ar", "Finance", "Riyadh"),
    ("stc pay", "talent@stcpay-example.com", "en", "Finance", "Riyadh"),
    ("Tamara", "jobs@tamara-example.com", "en", "Finance", "Riyadh"),
    # --- Engineering ---
    ("Saudi Aramco", "careers@aramco-example.com", "en", "Engineering", "Dhahran"),
    ("SABIC", "jobs@sabic-example.com", "en", "Engineering", "Riyadh"),
    ("ACWA Power", "jobs@acwapower-example.com", "en", "Engineering", "Riyadh"),
    ("Alfanar", "hr@alfanar-example.com", "en", "Engineering", "Riyadh"),
    # --- Legal ---
    ("Al-Bassam Law Firm", "info@bassamlaw-example.com", "ar", "Legal", "Riyadh"),
    ("Khoshaim & Associates", "careers@khoshaim-example.com", "en", "Legal", "Riyadh"),
    # --- Retail ---
    ("Jarir Bookstore", "careers@jarir-example.com", "ar", "Retail", "Riyadh"),
    ("Panda Retail", "jobs@panda-example.com", "ar", "Retail", "Jeddah"),
    ("Extra", "careers@extra-example.com", "ar", "Retail", "Khobar"),
    # --- Education ---
    ("Alfaisal University", "jobs@alfaisal-example.com", "en", "Education", "Riyadh"),
    ("Manarat Schools", "hr@manarat-example.com", "ar", "Education", "Jeddah"),
    # --- HR ---
    ("Bayt Recruitment", "careers@bayt-example.com", "en", "HR", "Riyadh"),
    ("Mihnati", "hr@mihnati-example.com", "ar", "HR", "Riyadh"),
    # --- Hospitality ---
    ("Ritz-Carlton Riyadh", "careers@ritz-example.com", "en", "Hospitality", "Riyadh"),
    ("Movenpick Hotels KSA", "hr@movenpick-example.com", "en", "Hospitality", "Jeddah"),
    # --- Logistics ---
    ("Aramex KSA", "jobs@aramex-example.com", "en", "Logistics", "Riyadh"),
    ("SMSA Express", "hr@smsa-example.com", "ar", "Logistics", "Dammam"),
]

SECTORS_NOTE = "بيانات تجريبية للعرض فقط — تُستبدل/تُوسّع حسب الاتفاق."


def build_db():
    if os.path.exists(DB):
        os.remove(DB)
    conn = sqlite3.connect(DB)
    conn.execute("""
        CREATE TABLE companies (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            company_name TEXT NOT NULL,
            email        TEXT NOT NULL UNIQUE,
            language     TEXT DEFAULT 'en',
            sector       TEXT,
            city         TEXT,
            source       TEXT DEFAULT 'sample',
            added_at     TEXT DEFAULT (datetime('now'))
        );
    """)
    conn.executemany(
        "INSERT INTO companies(company_name,email,language,sector,city) VALUES (?,?,?,?,?)",
        COMPANIES,
    )
    conn.commit()
    n = conn.execute("SELECT COUNT(*) FROM companies").fetchone()[0]
    # quick demo queries
    by_sector = conn.execute(
        "SELECT sector, COUNT(*) c FROM companies GROUP BY sector ORDER BY c DESC LIMIT 5"
    ).fetchall()
    conn.close()
    return n, by_sector


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    headers = ["company_name", "email", "language", "sector", "city"]
    ws.append(headers)
    # style header
    fill = PatternFill("solid", fgColor="2563EB")
    for col, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    for row in COMPANIES:
        ws.append(list(row))
    # column widths
    widths = [26, 38, 10, 16, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(XLSX)
    return len(COMPANIES)


if __name__ == "__main__":
    n, by_sector = build_db()
    rows = build_xlsx()
    print(f"✅ قاعدة البيانات: {DB}")
    print(f"   عدد الشركات: {n}")
    print("   توزيع حسب القطاع (أعلى 5):")
    for sector, c in by_sector:
        print(f"     - {sector}: {c}")
    print(f"\n✅ ملف Excel جاهز للبوت: {XLSX}")
    print(f"   عدد الصفوف: {rows}")
    print(f"\nℹ️  {SECTORS_NOTE}")
